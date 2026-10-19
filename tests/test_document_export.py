"""Tests for local XLSX and PDF report generation."""
import tempfile
import unittest
from pathlib import Path

from kasirkita.services.document_export import (
    export_sales_report_pdf, export_sales_report_xlsx,
)


class DocumentExportTests(unittest.TestCase):
    def setUp(self):
        self.report = {
            "start_date": "2026-10-10", "end_date": "2026-10-11",
            "sales_count": 2, "total_rupiah": 12500, "tax_rupiah": 0,
            "items_sold": 5, "average_sale_rupiah": 6250,
            "daily": [
                {"date": "2026-10-10", "sales_count": 1, "total_rupiah": 5000},
                {"date": "2026-10-11", "sales_count": 1, "total_rupiah": 7500},
            ],
        }

    def test_xlsx_contains_summary_and_daily_sheets(self):
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as directory:
            target = Path(directory) / "report.xlsx"
            export_sales_report_xlsx(self.report, target)
            workbook = load_workbook(str(target), read_only=True, data_only=True)
            try:
                self.assertEqual(workbook.sheetnames, ["Ringkasan", "Harian"])
                self.assertEqual(workbook["Harian"]["A2"].value, "2026-10-10")
                self.assertEqual(workbook["Ringkasan"]["B6"].value, 2)
            finally:
                workbook.close()

    def test_pdf_is_created(self):
        with tempfile.TemporaryDirectory() as directory:
            target = Path(directory) / "report.pdf"
            export_sales_report_pdf(self.report, target)
            self.assertTrue(target.is_file())
            self.assertTrue(target.read_bytes().startswith(b"%PDF-"))

    def test_invalid_report_rejected(self):
        invalid = dict(self.report, sales_count=True)
        with tempfile.TemporaryDirectory() as directory:
            with self.assertRaises(Exception):
                export_sales_report_xlsx(invalid, Path(directory) / "report.xlsx")


if __name__ == "__main__":
    unittest.main()
