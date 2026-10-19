"""Offline XLSX/PDF document exports for KasirKita.

Dependencies are optional at import time so the core POS can still start if an
export library is unavailable. Export actions return a clear error in that case.
"""
from __future__ import annotations

import os
from pathlib import Path
from typing import Mapping, Sequence


class DocumentExportError(RuntimeError):
    """Raised when a document cannot be validated or written."""


def _integer(value: object, label: str) -> int:
    if isinstance(value, bool) or not isinstance(value, int) or value < 0:
        raise DocumentExportError("Nilai tidak valid: " + label + ".")
    return value


def export_sales_report_xlsx(report: Mapping[str, object], destination: object) -> Path:
    """Write summary and daily detail into a formatted .xlsx workbook."""
    if not isinstance(report, Mapping):
        raise DocumentExportError("Data laporan tidak valid.")
    required = ("start_date", "end_date", "sales_count", "total_rupiah",
                "tax_rupiah", "items_sold", "average_sale_rupiah", "daily")
    if any(key not in report for key in required):
        raise DocumentExportError("Data laporan tidak lengkap.")
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise DocumentExportError(
            "Fitur Excel memerlukan paket openpyxl. Pasang dependensi build yang disediakan."
        ) from exc

    daily = report["daily"]
    if not isinstance(daily, (list, tuple)):
        raise DocumentExportError("Rincian harian laporan tidak valid.")
    target = _destination(destination, ".xlsx")
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Ringkasan"
    summary.append(["Laporan Penjualan KasirKita"])
    summary.append(["Tanggal awal", str(report["start_date"])])
    summary.append(["Tanggal akhir", str(report["end_date"])])
    summary.append([])
    summary.append(["Metrik", "Nilai"])
    for row in (
        ("Jumlah transaksi", _integer(report["sales_count"], "sales_count")),
        ("Total penjualan (rupiah)", _integer(report["total_rupiah"], "total_rupiah")),
        ("Total pajak tercatat (rupiah)", _integer(report["tax_rupiah"], "tax_rupiah")),
        ("Unit terjual", _integer(report["items_sold"], "items_sold")),
        ("Rata-rata transaksi (rupiah)", _integer(report["average_sale_rupiah"], "average_sale_rupiah")),
    ):
        summary.append(list(row))
    detail = workbook.create_sheet("Harian")
    detail.append(["Tanggal lokal", "Jumlah transaksi", "Total penjualan (rupiah)"])
    for item in daily:
        if not isinstance(item, Mapping) or not isinstance(item.get("date"), str):
            raise DocumentExportError("Baris laporan harian tidak valid.")
        detail.append([
            item["date"],
            _integer(item.get("sales_count"), "daily.sales_count"),
            _integer(item.get("total_rupiah"), "daily.total_rupiah"),
        ])
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="12304A")
            cell.alignment = Alignment(vertical="center")
        for column_cells in sheet.columns:
            max_width = max((len(str(cell.value or "")) for cell in column_cells), default=10)
            sheet.column_dimensions[get_column_letter(column_cells[0].column)].width = min(max_width + 3, 42)
    try:
        target.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(str(target))
    except (OSError, ValueError) as exc:
        raise DocumentExportError("File Excel gagal disimpan.") from exc
    finally:
        workbook.close()
    return target


def export_sales_report_pdf(report: Mapping[str, object], destination: object) -> Path:
    """Write a compact A4 sales report PDF without contacting external services."""
    if not isinstance(report, Mapping):
        raise DocumentExportError("Data laporan tidak valid.")
    required = ("start_date", "end_date", "sales_count", "total_rupiah",
                "tax_rupiah", "items_sold", "average_sale_rupiah", "daily")
    if any(key not in report for key in required):
        raise DocumentExportError("Data laporan tidak lengkap.")
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.lib.units import mm
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    except ImportError as exc:
        raise DocumentExportError(
            "Fitur PDF memerlukan paket reportlab. Pasang dependensi build yang disediakan."
        ) from exc

    daily = report["daily"]
    if not isinstance(daily, (list, tuple)):
        raise DocumentExportError("Rincian harian laporan tidak valid.")
    target = _destination(destination, ".pdf")
    styles = getSampleStyleSheet()
    story = [
        Paragraph("Laporan Penjualan KasirKita", styles["Title"]),
        Paragraph("Periode: " + str(report["start_date"]) + " sampai " + str(report["end_date"]), styles["Normal"]),
        Spacer(1, 6 * mm),
    ]
    summary_rows = [["Metrik", "Nilai"]]
    labels = [
        ("Jumlah transaksi", "sales_count"),
        ("Total penjualan (rupiah)", "total_rupiah"),
        ("Pajak tercatat (rupiah)", "tax_rupiah"),
        ("Unit terjual", "items_sold"),
        ("Rata-rata transaksi (rupiah)", "average_sale_rupiah"),
    ]
    for label, key in labels:
        summary_rows.append([label, str(_integer(report[key], key))])
    story.append(Table(summary_rows, colWidths=[85 * mm, 55 * mm], repeatRows=1, style=TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#12304A")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#CCD5DE")),
        ("PADDING", (0, 0), (-1, -1), 6),
    ])))
    story.append(Spacer(1, 7 * mm))
    rows = [["Tanggal lokal", "Transaksi", "Total penjualan (rupiah)"]]
    for item in daily:
        if not isinstance(item, Mapping) or not isinstance(item.get("date"), str):
            raise DocumentExportError("Baris laporan harian tidak valid.")
        rows.append([item["date"], str(_integer(item.get("sales_count"), "daily.sales_count")),
                     str(_integer(item.get("total_rupiah"), "daily.total_rupiah"))])
    story.append(Paragraph("Rincian Harian", styles["Heading2"]))
    story.append(Table(rows, colWidths=[50 * mm, 35 * mm, 60 * mm], repeatRows=1, style=TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#12304A")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#CCD5DE")),
        ("PADDING", (0, 0), (-1, -1), 5),
    ])))
    try:
        target.parent.mkdir(parents=True, exist_ok=True)
        SimpleDocTemplate(str(target), pagesize=landscape(A4), rightMargin=14*mm,
                          leftMargin=14*mm, topMargin=14*mm, bottomMargin=14*mm).build(story)
    except (OSError, ValueError) as exc:
        raise DocumentExportError("File PDF gagal disimpan.") from exc
    return target


def _destination(value: object, extension: str) -> Path:
    if not isinstance(value, (str, os.PathLike)) or not os.fspath(value).strip():
        raise DocumentExportError("Lokasi file tujuan tidak valid.")
    target = Path(value).expanduser()
    if target.suffix.lower() != extension:
        target = target.with_suffix(extension)
    return target


__all__ = ["DocumentExportError", "export_sales_report_pdf", "export_sales_report_xlsx"]
